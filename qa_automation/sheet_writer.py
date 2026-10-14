"""sheet_writer.py — เขียนผลลัพธ์ QA ลงชีท "AI Submit MDM" ผ่าน Google Sheets API
(gspread + service account) แทนการเปิด browser จริงแล้ว paste ผ่าน clipboard แบบเดิม —
เร็วกว่า, ไม่พึ่ง UI ของ Google Sheets ที่เปลี่ยนได้ตลอด, รันแบบ headless ได้เต็มรูปแบบ,
ไม่ต้อง login Google ผ่านเบราว์เซอร์อีกต่อไป.

Setup ที่ต้องทำครั้งเดียว (ดู README):
  1. สร้าง Service Account ใน Google Cloud Console แล้วเปิดใช้ Google Sheets API
  2. โหลดไฟล์ credentials JSON เก็บไว้ — ตั้ง path ไว้ที่ .env: GOOGLE_SERVICE_ACCOUNT_FILE
  3. แชร์สิทธิ์ "แก้ไขได้" ของชีท "AI Submit MDM" ให้ email ของ service account
     (รูปแบบ xxx@xxx.iam.gserviceaccount.com — ดูได้ในไฟล์ credentials JSON เอง)

Sheet: QA_SPREADSHEET_ID ใน .env
ไฟล์เดียวกับระบบ Helpdesk เดิม แต่เขียนลง 3 แท็บใหม่ที่เพิ่มมาสำหรับ flow CRM
(แท็บเดิม Submit (IOS)/No Submit/Android ของระบบ Helpdesk ไม่แตะ — ปล่อยไว้เฉยๆ):

  tab "CRM Submit"    (gid 2034718905) — iOS + MATCH ครบทุกจุด. 27 คอลัมน์ —
                         เหมือนแท็บ "Submit (IOS)" เดิมทุกอย่าง (25 คอลัมน์ A-Y)
                         บวก 2 คอลัมน์ท้ายที่คนกรอกเอง (ห้ามแตะ): Audit Check, Remark
  tab "CRM No Submit" (gid 1341767051) — ทุกเคสที่ไม่ match (ทั้ง iOS/Android).
                         27 คอลัมน์ — เหมือนแท็บ "No Submit" เดิมทุกอย่าง (26 คอลัมน์
                         A-Z) บวก 1 คอลัมน์ท้ายที่คนกรอกเอง (ห้ามแตะ): แก้ไข
  tab "CRM Android"   (gid 1529130860) — Android + MATCH ครบทุกจุด. 20 คอลัมน์ —
                         เหมือนแท็บ "Android" เดิมทุกอย่าง ไม่มีคอลัมน์เพิ่ม

คอลัมน์ที่ field=None (คนกรอกเองทีหลัง) อยู่ท้ายสุดของแต่ละแท็บเสมอ (ยืนยันจาก
schema จริงที่สำรวจมา) — เขียนแค่ช่วงคอลัมน์จริงแล้วหยุด ไม่แตะคอลัมน์ท้ายๆ เหล่านั้นเลย.

เคส K2 (ไม่มีเลขที่ใบคำขอ ITOS) ไม่เขียนที่ไหนเลย — ข้ามตั้งแต่ main.py.
"""
from __future__ import annotations

import logging
from typing import Optional

import gspread
from google.oauth2.service_account import Credentials

from qa_automation.config import env
from qa_automation.models import CaseRecord

logger = logging.getLogger(__name__)

_SCOPES = ["https://www.googleapis.com/auth/spreadsheets"]

SPREADSHEET_ID = env("QA_SPREADSHEET_ID")
QA_SHEET_URL = f"https://docs.google.com/spreadsheets/d/{SPREADSHEET_ID}/edit"

SUBMIT_GID = 2034718905      # gid จริงของ tab "CRM Submit" (27 คอลัมน์)
NO_SUBMIT_GID = 1341767051   # gid จริงของ tab "CRM No Submit" (27 คอลัมน์)
ANDROID_GID = 1529130860     # gid จริงของ tab "CRM Android" (20 คอลัมน์)

_COMMON_COLUMNS: list[tuple[str, str]] = [
    ("audit_time",        "เวลาที่ตรวจ"),
    ("case_id",          "Case ID"),
    ("match",             "MATCH (yes/no)"),
    ("customer_name",    "ชื่อ-นามสกุล ลูกค้า"),
    ("branch_id_name",    "Branch ID+Name"),
]

_PRICE_MODEL_COLUMNS: list[tuple[str, str]] = [
    ("loan_app_number",  "เลขที่ใบคำขอ"),
    ("model_name",        "ชื่อรุ่น"),
    ("storage",            "ความจุ"),
    ("color",               "สี"),
    ("price",               "ราคา"),
    ("down_payment_net", "เงินดาวน์สุทธิ"),
    ("down_payment",      "เงินดาวน์ที่ลูกค้าจ่ายจริง"),
    ("trade_in",            "Trade-in"),
    ("discount_package",       "ส่วนลด package ที่ได้รับ"),
    ("discount_down_payment", "ส่วนลดเงินดาวน์"),
    ("discount_other",         "ส่วนลดอื่นๆ"),
    ("discount_device",        "ส่วนลดค่าเครื่อง"),
]

# แท็บ "CRM Submit" (iOS เท่านั้น) — 25 คอลัมน์ automation + 2 คอลัมน์ท้ายห้ามแตะ
SUBMIT_COLUMNS: list[tuple[Optional[str], str]] = [
    *_COMMON_COLUMNS,
    *_PRICE_MODEL_COLUMNS,
    ("serial_number",     "Serial Number"),
    ("serial_source",     "Check S/N"),
    ("imei",                "IMEI"),
    ("imei_source",       "Check IMEI"),
    ("mdm_found",          "MDM (yes/no)"),
    ("macbook_mdm",       "MACBOOK MDM (yes/no)"),
    ("mdm_case",             "MDM Case"),
    ("remark",              "Remark:"),
    (None,                    "Audit Check"),   # คนกรอกเอง — automation ห้ามแตะ
    (None,                    "Remark"),        # คนกรอกเอง — automation ห้ามแตะ
]

# แท็บ "CRM No Submit" (รวม iOS/Android ที่ไม่ match) — 26 คอลัมน์ automation + 1 คอลัมน์ท้ายห้ามแตะ
NO_SUBMIT_COLUMNS: list[tuple[Optional[str], str]] = [
    *_COMMON_COLUMNS,
    ("platform", "Android / IOS"),
    *_PRICE_MODEL_COLUMNS,
    ("serial_number",     "Serial Number"),
    ("serial_source",     "Check S/N"),
    ("imei",                "IMEI"),
    ("imei_source",       "Check IMEI"),
    ("mdm_found",          "MDM (yes/no)"),
    ("macbook_mdm",       "MACBOOK MDM (yes/no)"),
    ("mdm_case",             "MDM Case"),
    ("remark",              "Remark:"),
    (None,                    "แก้ไข"),         # คนกรอกเอง — automation ห้ามแตะ
]

# แท็บ "CRM Android" — 20 คอลัมน์ (เหมือนเดิมทุกอย่าง ไม่มีคอลัมน์เพิ่ม)
ANDROID_COLUMNS: list[tuple[Optional[str], str]] = [
    *_COMMON_COLUMNS,
    *_PRICE_MODEL_COLUMNS,
    ("imei",          "IMEI"),
    ("imei_source",  "Check IMEI"),
    ("remark",          "Remark:"),
]

_TAB_SPECS: dict[str, tuple[int, list[tuple[Optional[str], str]]]] = {
    "CRM Submit":    (SUBMIT_GID, SUBMIT_COLUMNS),
    "CRM No Submit": (NO_SUBMIT_GID, NO_SUBMIT_COLUMNS),
    "CRM Android":   (ANDROID_GID, ANDROID_COLUMNS),
}


def _get_cell_value(rec: CaseRecord, field: Optional[str]) -> str:
    if field is None:  # คอลัมน์ที่ automation ห้ามแตะ — ไม่มีค่าจาก rec
        return ""
    return str(getattr(rec, field, "") or "")


def _leading_real_columns(columns: list[tuple[Optional[str], str]]) -> list[tuple[str, str]]:
    """คืนเฉพาะคอลัมน์ที่ automation ดูแล (field ไม่ใช่ None) — สมมติว่าคอลัมน์
    field=None (คนกรอกเอง) อยู่ท้ายสุดของ list เสมอ (ยืนยันจาก schema จริงที่สำรวจมา
    — ไม่มีช่องว่างแทรกกลาง) ถ้าไม่ใช่แบบนั้นจริง ให้ fail ดังๆ แทนที่จะเขียนทับผิดคอลัมน์
    แบบเงียบๆ."""
    real = [(f, l) for f, l in columns if f is not None]
    trailing = columns[len(real):]
    assert all(f is None for f, _ in trailing), (
        "sheet_writer: field=None columns ต้องอยู่ท้ายสุดเท่านั้น (โครงสร้าง schema เปลี่ยนไป "
        "จากที่เคยสำรวจไว้ — ต้องแก้ logic การเขียนใหม่)"
    )
    return real


class SheetsWriter:
    """เปิด spreadsheet ค้างไว้ครั้งเดียว (service account) แล้วเขียนทีละเคสตลอดการรัน."""

    def __init__(self, credentials_file: str, spreadsheet_id: str = SPREADSHEET_ID) -> None:
        creds = Credentials.from_service_account_file(credentials_file, scopes=_SCOPES)
        self._client = gspread.authorize(creds)
        self._spreadsheet = self._client.open_by_key(spreadsheet_id)
        self._worksheets: dict[int, gspread.Worksheet] = {}

    def _worksheet(self, gid: int) -> gspread.Worksheet:
        if gid not in self._worksheets:
            self._worksheets[gid] = self._spreadsheet.get_worksheet_by_id(gid)
        return self._worksheets[gid]

    def case_already_recorded(self, case_id: str) -> bool:
        """เช็คว่าเคสนี้เขียนลงชีทไปแล้วหรือยัง (คอลัมน์ B "Case ID" ของทั้ง 3 แท็บ) —
        กันเขียนซ้ำเวลารันคนละรอบ/คนละ process แล้วมาเจอเคสเดิมอีก (เคสยังค้างเป็น "เคสใหม่"
        ใน 7Chat เพราะไม่มีใครกด "รับเคส" — เจอบั๊กจริง: เคสเดียวกันถูกบันทึกซ้ำหลายแถวใน
        CRM No Submit ตอนรันหลายรอบติดกัน). เช็คจากชีทเอง (ไม่ใช่ตัวแปรในโปรแกรม) เพราะยัง
        ถูกต้องแม้เป็นคนละ process/รันคนละรอบกัน."""
        for gid, _ in _TAB_SPECS.values():
            ws = self._worksheet(gid)
            if case_id in ws.col_values(2):  # คอลัมน์ B = "Case ID" ทุกแท็บ
                return True
        return False

    def write_case(self, rec: CaseRecord) -> None:
        """เขียนผลลัพธ์เคสเดียวลงชีททันที — เลือกแท็บเองจาก platform/match ของ rec.
        เรียกทันทีหลังตรวจแต่ละเคสเสร็จ กันข้อมูลหายทั้งหมดถ้า script พังกลางทาง."""
        if rec.platform == "IOS" and rec.match == "Yes":
            tab_name = "CRM Submit"
        elif rec.platform == "Android" and rec.match == "Yes":
            tab_name = "CRM Android"
        else:
            tab_name = "CRM No Submit"

        gid, columns = _TAB_SPECS[tab_name]
        real_columns = _leading_real_columns(columns)
        ws = self._worksheet(gid)

        next_row = len(ws.col_values(1)) + 1
        values = [_get_cell_value(rec, field) for field, _ in real_columns]

        start_a1 = gspread.utils.rowcol_to_a1(next_row, 1)
        end_a1 = gspread.utils.rowcol_to_a1(next_row, len(real_columns))
        ws.update(f"{start_a1}:{end_a1}", [values], value_input_option="USER_ENTERED")

        highlight_cols = [
            col for col, (field, _) in enumerate(real_columns, start=1) if field in rec.failed_fields
        ]
        if highlight_cols:
            ws.batch_format([
                {
                    "range": gspread.utils.rowcol_to_a1(next_row, col),
                    "format": {"backgroundColor": {"red": 1, "green": 1, "blue": 0}},
                }
                for col in highlight_cols
            ])

        logger.info("Wrote case %s to '%s' tab (gid=%d) row %d", rec.case_id, tab_name, gid, next_row)


def write_local_xlsx_report(
    submit_records: list[CaseRecord],
    android_records: list[CaseRecord],
    no_submit_records: list[CaseRecord],
    path,
) -> None:
    """สำรอง: เขียนออก Excel ในเครื่องด้วย openpyxl — 3 sheet (Submit/Android/No Submit)."""
    from pathlib import Path
    from openpyxl import Workbook

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws_submit = wb.active
    ws_submit.title = "Submit"
    _fill_worksheet(ws_submit, submit_records, SUBMIT_COLUMNS)

    ws_android = wb.create_sheet("Android")
    _fill_worksheet(ws_android, android_records, ANDROID_COLUMNS)

    ws_no_submit = wb.create_sheet("No Submit")
    _fill_worksheet(ws_no_submit, no_submit_records, NO_SUBMIT_COLUMNS)

    wb.save(path)
    logger.info("Saved local XLSX to %s", path)


def _fill_worksheet(ws, records: list[CaseRecord], columns: list[tuple[Optional[str], str]]) -> None:
    from openpyxl.styles import Font, PatternFill, Alignment

    headers = [label for _, label in columns]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center")

    for rec in records:
        row = [_get_cell_value(rec, field) for field, _ in columns]
        ws.append(row)
        row_idx = ws.max_row
        for col_idx, (field, _) in enumerate(columns, start=1):
            val = _get_cell_value(rec, field)
            cell = ws.cell(row=row_idx, column=col_idx)
            if val.lower() == "yes":
                cell.fill = PatternFill("solid", fgColor="C6EFCE")
            elif val.lower() == "no":
                cell.fill = PatternFill("solid", fgColor="FFC7CE")
            cell.alignment = Alignment(horizontal="center")

    for col in ws.columns:
        max_len = max((len(str(c.value or "")) for c in col), default=8)
        ws.column_dimensions[col[0].column_letter].width = min(50, max(12, max_len + 2))
